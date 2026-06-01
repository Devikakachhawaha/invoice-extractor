import customtkinter as ctk
from tkinter import filedialog, messagebox
from streamlit import text
from tkinterdnd2 import DND_FILES, TkinterDnD
import pdfplumber
import pandas as pd
import re
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
import threading



# ==========================================
# APP SETTINGS
# ==========================================

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

# ==========================================
# MAIN WINDOW
# ==========================================

root = TkinterDnD.Tk()

root.title("PDF Extractor")

root.geometry("1000x650")

selected_files = []

# ==========================================
# DRAG & DROP
# ==========================================

def drop_files(event):

    global selected_files

    files = root.tk.splitlist(event.data)

    pdf_files = [
        file for file in files
        if file.lower().endswith(".pdf")
    ]

    selected_files.extend(pdf_files)

    # Remove duplicates
    selected_files[:] = list(set(selected_files))

    update_file_count()

# ==========================================
# UPDATE FILE COUNT
# ==========================================

def update_file_count():

    file_count_label.configure(
        text=f"{len(selected_files)} PDF files selected"
    )

# ==========================================
# BROWSE FILES
# ==========================================

def browse_files():

    global selected_files

    files = filedialog.askopenfilenames(
        title="Select PDF Files",
        filetypes=[("PDF Files", "*.pdf")]
    )

    selected_files.extend(files)

    # Remove duplicates
    selected_files[:] = list(set(selected_files))

    update_file_count()

# ==========================================
# EXTRACT TEXT FROM PDF
# ==========================================

def extract_pdf_text(pdf_path):

    full_text = ""

    try:

        with pdfplumber.open(pdf_path) as pdf:

            for page in pdf.pages:

                text = page.extract_text()

                if text:
                    full_text += text + "\n"

    except Exception as e:

        print(f"PDF Error: {e}")

    return full_text

# ==========================================
# EXTRACT REQUIRED DATA
# ==========================================

def extract_data(text):

    data = {
        "Invoice No. (from Shipping Bill)": "",
        "Port Code (from Shipping Bill)": "",
        "Shipping Bill No. (from Shipping Bill)": "",
        "Shipping Bill Date (from Shipping Bill)": "",
        "Invoice Date (from Shipping Bill)": "",
        "TAXABLE VALUE": 0.0,
        "IGST TAX AMOUNT": 0.0,
        "TAXABLE VALUE (FOB+FREIGHT+INSURANCE)": 0.0,
        "LUT": ""
    }

    # ======================================
    # CLEAN TEXT
    # ======================================

    clean_text = re.sub(r'\s+', ' ', text)

    # ======================================
    # PORT CODE / SB NO / SB DATE
    # ======================================

    sb_match = re.search(
        r'INDIAN CUSTOMS EDI SYSTEM\s+([A-Z0-9]+)\s+(\d+)\s+([0-9A-Z\-]+)',
        clean_text
    )

    if sb_match:

        data["Port Code (from Shipping Bill)"] = sb_match.group(1)

        data["Shipping Bill No. (from Shipping Bill)"] = sb_match.group(2)

        data["Shipping Bill Date (from Shipping Bill)"] = sb_match.group(3)

    # ======================================
    # INVOICE NUMBER + DATE
    # ======================================

    inv_match = re.search(
        r'(JTIPL/\d{4}/\d{3})\s+(\d{2}/\d{2}/\d{4})',
        clean_text
    )

    if inv_match:

        data["Invoice No. (from Shipping Bill)"] = inv_match.group(1)

        original_date = inv_match.group(2)

        converted_date = datetime.strptime(
            original_date,
            "%d/%m/%Y"
        ).strftime("%d-%b-%y").upper()

        data["Invoice Date (from Shipping Bill)"] = converted_date

    # ======================================
    # FOB VALUE
    # ======================================

    fob_match = re.search(
        r'LM\s+([0-9.]+)',
        clean_text
    )

    fob_value = 0.0

    if fob_match:

        fob_value = float(fob_match.group(1))

        data["TAXABLE VALUE"] = round(
            fob_value,
            2
        )

    # ======================================
    # FREIGHT VALUE
    # ======================================

    freight_match = re.search(
        r'FREIGHT\s+([0-9.]+)',
        clean_text,
        re.IGNORECASE
    )

    freight_value = 0.0

    if freight_match:

        freight_value = float(
            freight_match.group(1)
        )

    # ======================================
    # INSURANCE VALUE
    # ======================================

    insurance_match = re.search(
        r'INSURANCE\s+([0-9.]+)',
        clean_text,
        re.IGNORECASE
    )

    insurance_value = 0.0

    if insurance_match:

        insurance_value = float(
            insurance_match.group(1)
        )

    # ======================================
    # TOTAL TAXABLE VALUE
    # FOB + FREIGHT + INSURANCE
    # ======================================

    total_taxable = (
        fob_value +
        freight_value +
        insurance_value
    )

    data[
        "TAXABLE VALUE (FOB+FREIGHT+INSURANCE)"
    ] = round(
        total_taxable,
        2
    )

    # ======================================
    # IGST AMOUNT
    # ======================================

    igst_amt_match = re.search(
        r'3\.CESS AMT\s+([0-9.]+)\s+([0-9.]+)',
        clean_text
    )

    if igst_amt_match:

        data["IGST TAX AMOUNT"] = round(
            float(
                igst_amt_match.group(2)
            ),
            2
        )

    # ======================================
    # LUT VALUE
    # ======================================

    

    lut_match = re.search(
    r'1\.MODE\s+2\.ASSESS\s+3\.EXMN\s+4\.JOBBING\s+5\.MEIS\s+6\.DBK\s+7\.RODTP\s+8\.LICENCE\s+9\.DFRC\s+10\.RE-EXP\s+11\.LUT\s+'
    r'(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)',
    text  # use original text, not clean_text
    )

    if lut_match:
    # Groups: 1=MODE, 2=ASSESS, 3=EXMN, 4=JOBBING, 5=MEIS, 6=DBK, 7=RODTP, 8=LICENCE, 9=DFRC, 10=RE-EXP, 11=LUT
     data["LUT"] = lut_match.group(11).upper()

    return data

# ==========================================
# PROCESS FILES
# ==========================================

def process_files():

    if not selected_files:

        messagebox.showwarning(
            "Warning",
            "Please select PDF files"
        )

        return

    extract_button.configure(state="disabled")

    results = []

    duplicate_checker = set()

    total_files = len(selected_files)

    for index, file in enumerate(selected_files):

        try:

            # ==================================
            # UPDATE PROGRESS BAR
            # ==================================

            progress = (
                (index + 1) / total_files
            )

            progress_bar.set(progress)

            status_label.configure(
                text=f"Processing: {os.path.basename(file)}"
            )

            root.update_idletasks()

            # ==================================
            # EXTRACT TEXT
            # ==================================

            text = extract_pdf_text(file)

            # ==================================
            # SAVE DEBUG FILE
            # ==================================

            with open(
                "debug.txt",
                "w",
                encoding="utf-8"
            ) as f:

                f.write(text)

            # ==================================
            # EXTRACT DATA
            # ==================================

            extracted = extract_data(text)

            invoice_no = extracted[
                "Invoice No. (from Shipping Bill)"
            ]

            # ==================================
            # SKIP DUPLICATES
            # ==================================

            if invoice_no in duplicate_checker and invoice_no != "":

                continue

            duplicate_checker.add(invoice_no)

            results.append(extracted)

        except Exception as e:

            print(f"Processing Error: {e}")

    # ======================================
    # CREATE DATAFRAME
    # ======================================

    columns = [
        "S NO.",
        "Invoice No. (from Shipping Bill)",
        "Port Code (from Shipping Bill)",
        "Shipping Bill No. (from Shipping Bill)",
        "Shipping Bill Date (from Shipping Bill)",
        "Invoice Date (from Shipping Bill)",
        "TAXABLE VALUE",
        "IGST TAX AMOUNT",
        "TAXABLE VALUE (FOB+FREIGHT+INSURANCE)",
        "LUT"
    ]

    df = pd.DataFrame(results)

    # ======================================
    # SORT BY INVOICE DATE
    # ======================================

    df["SORT_DATE"] = pd.to_datetime(
        df["Invoice Date (from Shipping Bill)"],
        format="%d-%b-%y",
        errors="coerce"
    )

    df = df.sort_values(
        by="SORT_DATE",
        ascending=True
    )

    df = df.drop(columns=["SORT_DATE"])

    # ======================================
    # ADD SERIAL NUMBER
    # ======================================

    df.insert(
        0,
        "S NO.",
        range(1, len(df) + 1)
    )

    df = df[columns]

    # ======================================
    # SAVE EXCEL
    # ======================================

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_file = (
        f"Shipping_Bill_Data_{timestamp}.xlsx"
    )

    df.to_excel(output_file, index=False)

    # ======================================
    # FORMAT EXCEL
    # ======================================

    wb = load_workbook(output_file)

    ws = wb.active
    

    # ======================================
    # HEADER FORMATTING
    # ======================================

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FF0000"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ======================================
    # CENTER ALIGNMENT FOR ALL CELLS
    # ======================================

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    # ======================================
    # AUTO COLUMN WIDTH
    # ======================================

    for column in ws.columns:

        max_length = 0

        column_letter = column[0].column_letter

        for cell in column:

            try:

                if len(str(cell.value)) > max_length:

                    max_length = len(str(cell.value))

            except:
                pass

        adjusted_width = max_length + 5

        ws.column_dimensions[
            column_letter
        ].width = adjusted_width

    wb.save(output_file)

    # ======================================
    # COMPLETE
    # ======================================

    progress_bar.set(1)

    status_label.configure(
        text="Extraction Completed"
    )

    extract_button.configure(state="normal")

    messagebox.showinfo(
        "Success",
        f"Excel File Created Successfully!\n\n{output_file}"
    )

    os.startfile(output_file)

# ==========================================
# THREADING
# ==========================================

def start_extraction():

    threading.Thread(
        target=process_files
    ).start()

# ==========================================
# UI DESIGN
# ==========================================

title_label = ctk.CTkLabel(
    root,
    text="PDF Extractor",
    font=("Arial", 30, "bold")
)

title_label.pack(pady=25)

# ==========================================
# DRAG & DROP AREA
# ==========================================

drop_frame = ctk.CTkFrame(
    root,
    width=750,
    height=140
)

drop_frame.pack(pady=20)

drop_label = ctk.CTkLabel(
    drop_frame,
    text="Drag & Drop PDF Files Here",
    font=("Arial", 22)
)

drop_label.place(
    relx=0.5,
    rely=0.5,
    anchor="center"
)

drop_frame.drop_target_register(DND_FILES)

drop_frame.dnd_bind(
    "<<Drop>>",
    drop_files
)

# ==========================================
# BROWSE BUTTON
# ==========================================

browse_button = ctk.CTkButton(
    root,
    text="Browse PDF Files",
    command=browse_files,
    width=250,
    height=50,
    font=("Arial", 18)
)

browse_button.pack(pady=10)

# ==========================================
# FILE COUNT
# ==========================================

file_count_label = ctk.CTkLabel(
    root,
    text="0 PDF files selected",
    font=("Arial", 16)
)

file_count_label.pack()

# ==========================================
# PROGRESS BAR
# ==========================================

progress_bar = ctk.CTkProgressBar(
    root,
    width=700,
    height=20
)

progress_bar.pack(pady=25)

progress_bar.set(0)

# ==========================================
# STATUS LABEL
# ==========================================

status_label = ctk.CTkLabel(
    root,
    text="Waiting for PDF files...",
    font=("Arial", 15)
)

status_label.pack()

# ==========================================
# EXTRACT BUTTON
# ==========================================

extract_button = ctk.CTkButton(
    root,
    text="Extract Data",
    command=start_extraction,
    width=300,
    height=60,
    font=("Arial", 22, "bold")
)

extract_button.pack(pady=35)

# ==========================================
# RUN APP
# ==========================================

root.mainloop()